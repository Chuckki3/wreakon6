from __future__ import annotations

from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterator

from openpyxl import load_workbook

from wreakon6.domain.transaction import (
    CanonicalTransaction,
    TransactionDirection,
    TransactionStatus,
    TransactionType,
)


REQUIRED_COLUMNS = {
    "id",
    "walletNumber",
    "transactionDate",
    "transactionId",
    "partTransactionSerialNumber",
    "partTransactionType",
    "transactionAmount",
    "transactionCurrency",
    "narration",
    "narration2",
    "runningBalance",
    "valueDate",
    "isReversed",
    "reversalDate",
    "transactionType",
    "walletName",
    "status",
    "enteredBy",
    "createdAt",
    "postedBy",
    "thirdPartyTransactionReference",
    "updatedAt",
    "formattedAmount",
    "transactionReference",
}


def _parse_datetime(value) -> datetime:
    if isinstance(value, datetime):
        return value

    if value is None:
        raise ValueError("Missing datetime")

    return datetime.fromisoformat(str(value).strip())


def _parse_amount(value) -> Decimal:
    try:
        amount = Decimal(str(value))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"Invalid transaction amount: {value!r}") from exc

    if amount < 0:
        raise ValueError(f"Expected non-negative source amount, got {amount}")

    return amount


def _direction(value: str) -> TransactionDirection:
    value = (value or "").strip().upper()

    if value == "D":
        return TransactionDirection.DEBIT

    if value == "C":
        return TransactionDirection.CREDIT

    return TransactionDirection.UNKNOWN


def _transaction_type(
    source_type: str,
    narration: str | None,
) -> TransactionType:
    source_type = (source_type or "").strip().upper()
    narration_upper = (narration or "").strip().upper()

    if "HOLD ONLY" in narration_upper:
        return TransactionType.HOLD

    if "REFUND" in source_type:
        return TransactionType.REFUND

    if "AUTHORIZATION" in source_type:
        return TransactionType.AUTHORIZATION

    if "CREATION" in source_type:
        return TransactionType.CARD_CREATION

    return TransactionType.OTHER


def _status(value: str, is_reversed: bool) -> TransactionStatus:
    if is_reversed:
        return TransactionStatus.REVERSED

    value = (value or "").strip().upper()

    if value in {"P", "POSTED", "SUCCESS"}:
        return TransactionStatus.POSTED

    if value == "PENDING":
        return TransactionStatus.PENDING

    if value == "FAILED":
        return TransactionStatus.FAILED

    return TransactionStatus.UNKNOWN


def read_bank_xlsx(path: str | Path) -> Iterator[CanonicalTransaction]:
    """
    Stream the bank workbook in read-only mode.
    """

    path = Path(path)

    workbook = load_workbook(
        filename=path,
        read_only=True,
        data_only=True,
    )

    try:
        if "data" not in workbook.sheetnames:
            raise ValueError(
                f"Bank workbook must contain a 'data' sheet. "
                f"Found: {workbook.sheetnames}"
            )

        worksheet = workbook["data"]

        rows = worksheet.iter_rows(values_only=True)

        try:
            headers = next(rows)
        except StopIteration as exc:
            raise ValueError("Bank workbook is empty") from exc

        columns = {
            str(column).strip()
            for column in headers
            if column is not None
        }

        missing = REQUIRED_COLUMNS - columns

        if missing:
            raise ValueError(
                f"Bank workbook is missing required columns: {sorted(missing)}"
            )

        header_index = {
            str(column).strip(): index
            for index, column in enumerate(headers)
        }

        for row_number, row in enumerate(rows, start=2):
            record = {
                column: row[index]
                for column, index in header_index.items()
            }

            try:
                narration = record["narration"]
                transaction_type = _transaction_type(
                    record["transactionType"],
                    narration,
                )

                is_reversed = bool(record["isReversed"])

                yield CanonicalTransaction(
                    source="BANK",
                    source_record_id=str(record["id"]),
                    transaction_id=(
                        str(record["transactionId"])
                        if record["transactionId"] is not None
                        else None
                    ),
                    reference=(
                        str(record["transactionReference"])
                        if record["transactionReference"] is not None
                        else None
                    ),
                    account_id=(
                        str(record["walletNumber"])
                        if record["walletNumber"] is not None
                        else None
                    ),
                    event_at=_parse_datetime(record["transactionDate"]),
                    value_at=_parse_datetime(record["valueDate"]),
                    amount=_parse_amount(record["transactionAmount"]),
                    currency=str(record["transactionCurrency"]).strip().upper(),
                    direction=_direction(record["partTransactionType"]),
                    transaction_type=transaction_type,
                    status=_status(record["status"], is_reversed),
                    is_hold=transaction_type == TransactionType.HOLD,
                    is_reversal=is_reversed,
                    counterparty_reference=(
                        str(record["thirdPartyTransactionReference"])
                        if record["thirdPartyTransactionReference"] is not None
                        else None
                    ),
                    narration=(
                        str(narration).strip()
                        if narration is not None
                        else None
                    ),
                )

            except Exception as exc:
                raise ValueError(
                    f"Invalid bank record at worksheet row {row_number}: {exc}"
                ) from exc

    finally:
        workbook.close()